"""文档摄取与幂等索引测试。"""

import hashlib
from pathlib import Path
from typing import cast

import pytest
from openpyxl import Workbook  # type: ignore[import-untyped]
from pptx import Presentation
from agentscope.message import TextBlock
from agentscope.rag import (
    ApproxTokenChunker,
    Chunk,
    DocumentSummary,
    ExcelParser,
    KnowledgeBase,
    PPTParser,
    TextParser,
)

from chatbot_rag.rag import (
    ContextPreservingChunker,
    DocumentIngestor,
    IngestionStage,
    MediaAssetStore,
)
from chatbot_rag.rag.media_assets import extract_media_asset_ids
from chatbot_rag.rag.document_ingestor import (
    INGESTION_PIPELINE_KEY,
    INGESTION_PIPELINE_VERSION,
)


class FakeKnowledgeBase:
    """记录文档写入和删除操作的知识库替身。"""

    def __init__(self, summaries: list[DocumentSummary] | None = None) -> None:
        """使用可选的已有文档摘要初始化。"""
        self.summaries = summaries or []
        self.inserted: list[
            tuple[list[Chunk], str | None, dict[str, object] | None]
        ] = []
        self.deleted: list[str] = []

    async def list_documents(self) -> list[DocumentSummary]:
        """返回预设的文档摘要。"""
        return self.summaries

    async def insert_document(
        self,
        chunks: list[Chunk],
        document_id: str | None = None,
        document_metadata: dict[str, object] | None = None,
    ) -> str:
        """记录索引写入并返回文档标识。"""
        self.inserted.append((chunks, document_id, document_metadata))
        return document_id or "generated"

    async def delete_document(self, document_id: str) -> None:
        """记录待删除的历史文档标识。"""
        self.deleted.append(document_id)


def _create_ingestor(knowledge_base: FakeKnowledgeBase) -> DocumentIngestor:
    """创建仅处理 Markdown 的测试摄取器。"""
    return DocumentIngestor(
        knowledge_base=cast(KnowledgeBase, knowledge_base),
        chunker=ApproxTokenChunker(chunk_size=32, overlap=4),
        parsers={".md": TextParser()},
    )


@pytest.mark.asyncio
async def test_ingest_directory_indexes_changed_document_then_deletes_old(
    tmp_path: Path,
) -> None:
    """内容变化时应先写入新版本，再清理旧版本。"""
    document_path = tmp_path / "guide.md"
    document_path.write_text("AgentScope native RAG", encoding="utf-8")
    knowledge_base = FakeKnowledgeBase(
        [
            DocumentSummary(
                document_id="old-version",
                source="guide.md",
                chunk_count=1,
                metadata={
                    "source_path": "guide.md",
                    "content_sha256": "old-hash",
                },
            ),
        ],
    )

    stages: list[IngestionStage] = []
    summary = await _create_ingestor(knowledge_base).ingest_directory(
        tmp_path,
        progress_callback=stages.append,
    )

    assert summary.indexed_documents == 1
    assert summary.skipped_documents == 0
    assert summary.deleted_documents == 1
    assert stages == [IngestionStage.PARSING, IngestionStage.INDEXING]
    assert knowledge_base.deleted == ["old-version"]
    chunks, document_id, metadata = knowledge_base.inserted[0]
    assert chunks[0].source == "guide.md"
    assert document_id is not None
    assert metadata is not None
    assert metadata["source_path"] == "guide.md"
    assert metadata["content_sha256"] == hashlib.sha256(
        document_path.read_bytes(),
    ).hexdigest()


@pytest.mark.asyncio
async def test_ingest_directory_skips_unchanged_document(
    tmp_path: Path,
) -> None:
    """内容摘要未变化时不应重复调用嵌入与写库。"""
    content = b"unchanged knowledge"
    (tmp_path / "guide.md").write_bytes(content)
    knowledge_base = FakeKnowledgeBase(
        [
            DocumentSummary(
                document_id="current-version",
                source="guide.md",
                chunk_count=1,
                metadata={
                    "source_path": "guide.md",
                    "content_sha256": hashlib.sha256(content).hexdigest(),
                    INGESTION_PIPELINE_KEY: INGESTION_PIPELINE_VERSION,
                },
            ),
        ],
    )

    summary = await _create_ingestor(knowledge_base).ingest_directory(tmp_path)

    assert summary.indexed_documents == 0
    assert summary.skipped_documents == 1
    assert summary.deleted_documents == 0
    assert knowledge_base.inserted == []


@pytest.mark.asyncio
async def test_ingest_directory_deletes_removed_document(
    tmp_path: Path,
) -> None:
    """文件从受管目录移除后应清理对应的知识库文档。"""
    knowledge_base = FakeKnowledgeBase(
        [
            DocumentSummary(
                document_id="removed-document",
                source="removed.md",
                chunk_count=1,
                metadata={
                    "source_path": "removed.md",
                    "content_sha256": "old-hash",
                },
            ),
        ],
    )

    summary = await _create_ingestor(knowledge_base).ingest_directory(tmp_path)

    assert summary.deleted_documents == 1
    assert knowledge_base.deleted == ["removed-document"]


@pytest.mark.asyncio
async def test_ingest_directory_deletes_old_index_for_empty_document(
    tmp_path: Path,
) -> None:
    """文档内容变为空时不应继续保留旧的可检索内容。"""
    (tmp_path / "empty.md").write_text("", encoding="utf-8")
    knowledge_base = FakeKnowledgeBase(
        [
            DocumentSummary(
                document_id="previous-content",
                source="empty.md",
                chunk_count=1,
                metadata={
                    "source_path": "empty.md",
                    "content_sha256": "old-hash",
                },
            ),
        ],
    )

    summary = await _create_ingestor(knowledge_base).ingest_directory(tmp_path)

    assert summary.skipped_documents == 1
    assert summary.deleted_documents == 1
    assert knowledge_base.deleted == ["previous-content"]


@pytest.mark.asyncio
async def test_ingest_directory_rejects_missing_directory(
    tmp_path: Path,
) -> None:
    """文档目录不存在时应返回明确错误。"""
    missing_path = tmp_path / "missing"

    with pytest.raises(FileNotFoundError, match="does not exist"):
        await _create_ingestor(FakeKnowledgeBase()).ingest_directory(
            missing_path,
        )


@pytest.mark.asyncio
async def test_ingest_directory_commits_retrievable_media_manifest(
    tmp_path: Path,
) -> None:
    """摄取完成后应同时提交索引媒体标记和文档图片清单。"""
    document_path = tmp_path / "guide.markdown"
    document_path.write_text(
        "打开设置。\n\n"
        "![设置页面](https://images.example.com/setting.png)",
        encoding="utf-8",
    )
    media_store = MediaAssetStore(
        tmp_path / "media",
        ("images.example.com",),
    )
    knowledge_base = FakeKnowledgeBase()
    ingestor = DocumentIngestor(
        knowledge_base=cast(KnowledgeBase, knowledge_base),
        chunker=ApproxTokenChunker(chunk_size=32, overlap=4),
        media_store=media_store,
    )

    summary = await ingestor.ingest_directory(tmp_path)

    assert summary.indexed_documents == 1
    chunks, _, metadata = knowledge_base.inserted[0]
    asset_ids = [
        asset_id
        for chunk in chunks
        if isinstance(chunk.content, TextBlock)
        for asset_id in extract_media_asset_ids(chunk.content.text)
    ]
    assert len(asset_ids) == 1
    assert media_store.has_document("guide.markdown") is True
    assert metadata is not None
    assert metadata[INGESTION_PIPELINE_KEY] == INGESTION_PIPELINE_VERSION


def test_default_parsers_register_expected_document_extensions(
    tmp_path: Path,
) -> None:
    """默认摄取器应登记全部已承诺支持的文档扩展名。"""
    ingestor = DocumentIngestor(
        knowledge_base=cast(KnowledgeBase, FakeKnowledgeBase()),
        chunker=ApproxTokenChunker(),
    )
    expected_extensions = {
        ".docx",
        ".md",
        ".markdown",
        ".pdf",
        ".pptx",
        ".txt",
        ".xls",
        ".xlsx",
    }
    for extension in expected_extensions:
        (tmp_path / f"document{extension}").write_bytes(b"placeholder")
    (tmp_path / "legacy.doc").write_bytes(b"placeholder")
    (tmp_path / "legacy.ppt").write_bytes(b"placeholder")

    supported_files = ingestor._iter_supported_files(tmp_path)

    assert {path.suffix.lower() for path in supported_files} == (
        expected_extensions
    )
    powerpoint_parser = cast(PPTParser, ingestor._parsers[".pptx"])
    excel_parser = cast(ExcelParser, ingestor._parsers[".xlsx"])
    assert powerpoint_parser.include_image is False
    assert excel_parser.include_image is False
    assert excel_parser.separate_sheet is True
    assert ingestor._parsers[".xls"] is excel_parser


@pytest.mark.asyncio
async def test_default_parsers_ingest_txt_pptx_and_xlsx(
    tmp_path: Path,
) -> None:
    """TXT、PPTX 和 XLSX 应通过 AgentScope 原生 Parser 完成摄取。"""
    (tmp_path / "notes.txt").write_text("纯文本知识", encoding="utf-8")

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "部署说明"
    slide.placeholders[1].text = "服务需要先完成配置。"
    presentation.save(tmp_path / "deployment.pptx")

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "产品清单"
    worksheet.append(["产品", "状态"])
    worksheet.append(["知识助手", "可用"])
    workbook.save(tmp_path / "products.xlsx")
    workbook.close()

    knowledge_base = FakeKnowledgeBase()
    summary = await DocumentIngestor(
        knowledge_base=cast(KnowledgeBase, knowledge_base),
        chunker=ContextPreservingChunker(
            ApproxTokenChunker(chunk_size=128, overlap=16),
        ),
    ).ingest_directory(tmp_path)

    assert summary.indexed_documents == 3
    chunks_by_source = {
        cast(str, metadata["source_path"]): chunks
        for chunks, _, metadata in knowledge_base.inserted
        if metadata is not None
    }
    text_by_source = {
        source: "\n".join(
            chunk.content.text
            for chunk in chunks
            if isinstance(chunk.content, TextBlock)
        )
        for source, chunks in chunks_by_source.items()
    }
    assert "纯文本知识" in text_by_source["notes.txt"]
    assert "幻灯片：1" in text_by_source["deployment.pptx"]
    assert "部署说明" in text_by_source["deployment.pptx"]
    assert "工作表：产品清单" in text_by_source["products.xlsx"]
    assert "知识助手" in text_by_source["products.xlsx"]
