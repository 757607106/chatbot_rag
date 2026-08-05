"""媒体感知文档解析器测试。"""

from __future__ import annotations

import base64
from pathlib import Path

import pytest
from agentscope.message import TextBlock
from docx import Document

from chatbot_rag.rag import MediaAssetStore
from chatbot_rag.rag.media_assets import extract_media_asset_ids
from chatbot_rag.rag.media_parsers import (
    MarkdownMediaParser,
    PDFMediaParser,
    WordMediaParser,
)
from chatbot_rag.rag.contextual_chunker import RETRIEVAL_CONTEXT_KEY

PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR4"
    "nGNgYAAAAAMAASsJTYQAAAAASUVORK5CYII=",
)


@pytest.mark.asyncio
async def test_markdown_parser_replaces_remote_image_with_asset_reference(
    tmp_path: Path,
) -> None:
    """Markdown 外链图片应登记到媒体仓库并保留原始文本顺序。"""
    store = MediaAssetStore(tmp_path / "media", ("images.example.com",))
    parser = MarkdownMediaParser(store)

    sections = await parser.parse(
        (
            "第一步：打开设置。\n\n"
            "![设置页面](https://images.example.com/setting.png)\n\n"
            "第二步：保存。"
        ).encode(),
        "guide.md",
    )

    assert len(sections) == 1
    assert isinstance(sections[0].content, TextBlock)
    asset_ids = extract_media_asset_ids(sections[0].content.text)
    assert len(asset_ids) == 1
    assert "images.example.com" not in sections[0].content.text
    assert sections[0].content.text.index("第一步") < (
        sections[0].content.text.index("<chatbot-media")
    ) < sections[0].content.text.index("第二步")
    assert store.describe(asset_ids[0]).filename == "设置页面"


@pytest.mark.asyncio
async def test_markdown_parser_replaces_http_remote_image(
    tmp_path: Path,
) -> None:
    """HTTP 外链图片应与 HTTPS 一样登记到媒体仓库。"""
    store = MediaAssetStore(tmp_path / "media", ("images.example.com",))
    parser = MarkdownMediaParser(store)

    sections = await parser.parse(
        (
            "![设置页面](http://images.example.com/setting.png)\n\n"
            "第二步：保存。"
        ).encode(),
        "guide.md",
    )

    assert len(sections) == 1
    assert isinstance(sections[0].content, TextBlock)
    asset_ids = extract_media_asset_ids(sections[0].content.text)
    assert len(asset_ids) == 1
    assert "images.example.com" not in sections[0].content.text
    assert store.describe(asset_ids[0]).filename == "设置页面"


@pytest.mark.asyncio
async def test_markdown_parser_skips_disallowed_remote_image(
    tmp_path: Path,
) -> None:
    """不允许主机上的远程图片应被跳过，不影响文本索引。"""
    store = MediaAssetStore(tmp_path / "media", ("images.example.com",))
    parser = MarkdownMediaParser(store)

    sections = await parser.parse(
        (
            "第一步：打开设置。\n\n"
            "![允许图片](https://images.example.com/setting.png)\n\n"
            "![不允许图片](https://other.example.com/forbidden.png)\n\n"
            "第二步：保存。"
        ).encode(),
        "guide.md",
    )

    assert len(sections) == 1
    assert isinstance(sections[0].content, TextBlock)
    asset_ids = extract_media_asset_ids(sections[0].content.text)
    assert len(asset_ids) == 1
    assert "other.example.com" not in sections[0].content.text
    assert "forbidden.png" not in sections[0].content.text
    assert "第一步" in sections[0].content.text
    assert "第二步" in sections[0].content.text


@pytest.mark.asyncio
async def test_markdown_parser_preserves_full_heading_context(
    tmp_path: Path,
) -> None:
    """每个 Markdown 自然章节都应携带完整标题路径且不混合打印模式。"""
    store = MediaAssetStore(tmp_path / "media", ("images.example.com",))
    parser = MarkdownMediaParser(store)

    sections = await parser.parse(
        (
            "# 打印手册\n\n"
            "## 本地打印模式\n\n"
            "### 设置打印模板\n\n"
            "在打印管理器点击模板编辑。\n\n"
            "![本地入口](https://images.example.com/local.png)\n\n"
            "## Web打印模式\n\n"
            "### 设置打印模板\n\n"
            "点击打印按钮旁的自定义编辑。"
        ).encode(),
        "print-guide.md",
    )

    assert len(sections) == 2
    local_section, web_section = sections
    assert local_section.metadata[RETRIEVAL_CONTEXT_KEY] == (
        "# 打印手册\n## 本地打印模式\n### 设置打印模板"
    )
    assert web_section.metadata[RETRIEVAL_CONTEXT_KEY] == (
        "# 打印手册\n## Web打印模式\n### 设置打印模板"
    )
    assert isinstance(local_section.content, TextBlock)
    assert "Web打印模式" not in local_section.content.text
    assert len(extract_media_asset_ids(local_section.content.text)) == 1


@pytest.mark.asyncio
async def test_markdown_parser_ignores_headings_inside_code_fences(
    tmp_path: Path,
) -> None:
    """代码围栏中的井号文本不得被误判为 Markdown 标题。"""
    store = MediaAssetStore(tmp_path / "media", ())
    parser = MarkdownMediaParser(store)

    sections = await parser.parse(
        (
            "# 使用说明\n\n"
            "```text\n"
            "## 这不是标题\n"
            "```\n\n"
            "代码示例结束。"
        ).encode(),
        "guide.md",
    )

    assert len(sections) == 1
    assert sections[0].metadata[RETRIEVAL_CONTEXT_KEY] == "# 使用说明"
    assert isinstance(sections[0].content, TextBlock)
    assert "## 这不是标题" in sections[0].content.text


@pytest.mark.asyncio
async def test_markdown_parser_keeps_images_inside_code_fences_literal(
    tmp_path: Path,
) -> None:
    """代码示例中的图片语法不得登记成真实文档资产。"""
    store = MediaAssetStore(tmp_path / "media", ("images.example.com",))
    parser = MarkdownMediaParser(store)

    sections = await parser.parse(
        (
            "# 使用说明\n\n"
            "```markdown\n"
            "![示例](https://images.example.com/example.png)\n"
            "```\n"
        ).encode(),
        "guide.md",
    )

    assert len(sections) == 1
    assert isinstance(sections[0].content, TextBlock)
    assert "https://images.example.com/example.png" in sections[0].content.text
    assert not extract_media_asset_ids(sections[0].content.text)


@pytest.mark.asyncio
async def test_markdown_parser_replaces_skipped_level_heading_siblings(
    tmp_path: Path,
) -> None:
    """标题级别跳跃时，同级小节仍不得错误继承前一个兄弟标题。"""
    store = MediaAssetStore(tmp_path / "media", ())
    parser = MarkdownMediaParser(store)

    sections = await parser.parse(
        (
            "# 产品手册\n\n"
            "### 基础版\n\n"
            "基础版内容。\n\n"
            "### 专业版\n\n"
            "专业版内容。"
        ).encode(),
        "guide.md",
    )

    assert len(sections) == 2
    assert sections[0].metadata[RETRIEVAL_CONTEXT_KEY] == (
        "# 产品手册\n### 基础版"
    )
    assert sections[1].metadata[RETRIEVAL_CONTEXT_KEY] == (
        "# 产品手册\n### 专业版"
    )
    assert "基础版" not in sections[1].metadata[RETRIEVAL_CONTEXT_KEY]


@pytest.mark.asyncio
async def test_pdf_parser_extracts_page_images_to_media_references(
    tmp_path: Path,
) -> None:
    """PDF 页面图片应保存为本地资产并保留页码关联。"""
    source = Path("tests/docs_test/内容开发工程师岗位指导说明书.pdf")
    store = MediaAssetStore(tmp_path / "media", ())
    parser = PDFMediaParser(store)

    sections = await parser.parse(source.read_bytes(), source.name)

    referenced_ids = [
        asset_id
        for section in sections
        if isinstance(section.content, TextBlock)
        for asset_id in extract_media_asset_ids(section.content.text)
    ]
    assert len(referenced_ids) == 2
    assert {section.metadata["page"] for section in sections} >= {5, 6}
    for asset_id in referenced_ids:
        assert store.describe(asset_id).filename.endswith(".png")


@pytest.mark.asyncio
async def test_word_parser_extracts_embedded_image_next_to_text(
    tmp_path: Path,
) -> None:
    """Word 内嵌图片应保存，并关联到图片前的说明文本。"""
    image_path = tmp_path / "step.png"
    image_path.write_bytes(PNG_BYTES)
    document_path = tmp_path / "guide.docx"
    document = Document()
    document.add_paragraph("请打开系统设置。")
    document.add_picture(str(image_path))
    document.save(document_path)
    store = MediaAssetStore(tmp_path / "media", ())
    parser = WordMediaParser(store)

    sections = await parser.parse(document_path.read_bytes(), "guide.docx")

    assert len(sections) == 1
    assert isinstance(sections[0].content, TextBlock)
    assert "请打开系统设置" in sections[0].content.text
    asset_ids = extract_media_asset_ids(sections[0].content.text)
    assert len(asset_ids) == 1
    assert store.describe(asset_ids[0]).filename.startswith("guide-")
